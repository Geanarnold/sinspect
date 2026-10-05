# Gera app/catalogo.js a partir de catalogo/CATALOGO.xlsx e dos CSVs. Uso: python3 gerar_catalogo_js.py
import openpyxl,json,csv,datetime,os
os.chdir(os.path.dirname(os.path.abspath(__file__)))
wb=openpyxl.load_workbook('catalogo/CATALOGO.xlsx',data_only=True)
cat={'versao':str(datetime.date.today()),'colunas':{},'colunas_sa':{},'travessas':[],'diagonais':[],'uniao':[],'produtos':{},'composicao':[]}
for r in csv.DictReader(open('catalogo/colunas_v2.csv',encoding='utf-8')):
    cat['colunas'].setdefault(r['face_frontal_mm'],{})[r['espessura_mm']]=float(r['kg_por_m_calc_7.85'])
for r in csv.DictReader(open('catalogo/colunas_sa_altura.csv')):
    cat['colunas_sa'].setdefault(r['face_frontal_mm'],{})[r['altura_mm']]=r['sa']
def rows(name):
    ws=wb[name]; hdr=[c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2,values_only=True):
        d=dict(zip(hdr,row))
        if d.get(hdr[0]): yield d
for name,key in (('Travessas','travessas'),('Diagonais','diagonais')):
    for d in rows(name):
        if d['Status']=='EXCLUÍDO DA BUSCA' or not d['Código SA']: continue
        cat[key].append({'nome':d['Descrição'],'sa':d['Código SA'],'cc':d['Cota A: centro a centro (mm)'],'total':d['Cota B: comprimento total (mm)']})
for d in rows('Travessa união'):
    if isinstance(d['Comprimento total (mm)'],(int,float)):
        cat['uniao'].append({'nome':d['Descrição'],'col':d['Coluna'],'co':d['Código CO (conjunto)'],'peso':d['Peso (kg, desenho)'],'total':d['Comprimento total (mm)']})
for d in rows('Produtos'):
    cat['produtos'][d['ID']]={'tipo':d['Tipo'],'desc':d['Descrição'],'codigo':d['Código'] or '','peso':d['Peso usado no cálculo'] if isinstance(d['Peso usado no cálculo'],(int,float)) else None,'unid':d['Unidade'],'aco':d['Aço / material'] or '','dim':d.get('Dimensões (mm)') if isinstance(d.get('Dimensões (mm)'),(int,float)) else None}
for d in rows('Composição'):
    cat['composicao'].append({'pai':d['Conjunto (pai)'],'item':d['Item (ID em Produtos)'],'qtd':d['Quantidade'],'contagem':d['Contagem']})
open('app/catalogo.js','w',encoding='utf-8').write('// Gerado de ../catalogo/CATALOGO.xlsx por gerar_catalogo_js.py — não editar à mão.\nwindow.CATALOGO = '+json.dumps(cat,ensure_ascii=False,indent=1)+';\n')
print('ok',{k:len(v) for k,v in cat.items() if isinstance(v,(list,dict))})
